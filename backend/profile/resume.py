"""Resume upload (UP-6): validation, storage, text extraction and retention.

REQ-F1-08  PDF and DOCX only, otherwise "Only PDF and DOCX files are accepted."
REQ-F1-09  Larger than 5 MB: "File exceeds the 5 MB limit."
REQ-F1-10  Plain text is extracted from an accepted file.
REQ-F1-11  Under 100 characters of text: the resume is reported unreadable and
           manual skill entry is offered instead.
REQ-F1-18  The file is kept for at most 30 days after processing.
REQ-F1-38  If processing fails, a clear error and an unchanged profile.
NFR-SEC-06 Type and size are checked before anything is stored or parsed.

The type is decided from the file's contents, never from its name or the
client's declared content type. Files are stored outside the web root under
random names; the original name is only recorded.
"""

import io
import os
import tempfile
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.profile.models import ResumeUpload

MAX_BYTES = 5 * 1024 * 1024
MIN_TEXT_CHARS = 100
RETENTION = timedelta(days=30)
MAX_PDF_PAGES = 50
MAX_DOCX_EXPANDED_BYTES = 50 * 1024 * 1024  # zip bomb guard

TYPE_MESSAGE = "Only PDF and DOCX files are accepted."
SIZE_MESSAGE = "File exceeds the 5 MB limit."
UNREADABLE_MESSAGE = (
    "We could not read any text in this resume. You can add your skills manually instead."
)
FAILED_MESSAGE = "We could not process this resume. Your profile has not been changed."


class ResumeRejected(ValueError):
    """The upload is refused before anything is stored."""

    def __init__(self, message: str, status_code: int):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


class ExtractionFailed(Exception):
    """The file passed the type check but its text could not be extracted."""


def storage_dir() -> Path:
    """Outside the web root. RESUME_STORAGE_DIR in deployments; a private temp
    folder otherwise."""
    configured = os.environ.get("RESUME_STORAGE_DIR")
    path = Path(configured) if configured else Path(tempfile.gettempdir()) / "upskill-resumes"
    path.mkdir(mode=0o700, parents=True, exist_ok=True)
    return path


def detect_type(data: bytes) -> str | None:
    """'pdf' or 'docx' from the file's own bytes, or None."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def check_upload(data: bytes) -> str:
    """Size, then type (REQ-F1-09, F1-08); returns the detected type."""
    if len(data) > MAX_BYTES:
        raise ResumeRejected(SIZE_MESSAGE, 413)
    file_type = detect_type(data)
    if file_type is None:
        raise ResumeRejected(TYPE_MESSAGE, 415)
    return file_type


def extract_text(data: bytes, file_type: str) -> str:
    """Plain text of the resume (REQ-F1-10). Raises ExtractionFailed."""
    try:
        if file_type == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionFailed("encrypted PDF")
            pages = reader.pages[:MAX_PDF_PAGES]
            return "\n".join(page.extract_text() or "" for page in pages)

        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_EXPANDED_BYTES:
                raise ExtractionFailed("archive expands too far")
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except ExtractionFailed:
        raise
    except Exception as error:  # any parser failure is a failed upload, not a 500
        raise ExtractionFailed(type(error).__name__) from error


def clean_filename(name: str | None) -> str:
    """The client's file name for display only: no path, no control characters."""
    base = (name or "").replace("\\", "/").rsplit("/", 1)[-1]
    base = "".join(c for c in base if unicodedata.category(c)[0] != "C").strip()
    return (base or "resume")[:255]


@dataclass(frozen=True)
class UploadResult:
    upload: ResumeUpload
    message: str | None


class ResumeService:
    def __init__(self, session: Session, now=lambda: datetime.now(UTC)):
        self._session = session
        self._now = now

    def upload(self, student_id: int, filename: str | None, data: bytes) -> UploadResult:
        file_type = check_upload(data)  # nothing is stored if this raises

        stored = storage_dir() / f"{uuid.uuid4().hex}.{file_type}"
        stored.write_bytes(data)
        stored.chmod(0o600)

        try:
            text = extract_text(data, file_type)
            status = "extracted" if len(text.strip()) >= MIN_TEXT_CHARS else "unreadable"
        except ExtractionFailed:
            status = "failed"
        message = {"unreadable": UNREADABLE_MESSAGE, "failed": FAILED_MESSAGE}.get(status)

        processed = self._now()
        upload = ResumeUpload(
            student_id=student_id,
            original_filename=clean_filename(filename),
            stored_path=str(stored),
            file_type=file_type,
            size_bytes=len(data),
            status=status,
            processed_at=processed,
            delete_after=processed + RETENTION,
        )
        if status != "extracted":
            # Nothing will ever read an unusable file again: keep no copy.
            stored.unlink(missing_ok=True)
            upload.deleted_at = processed
        try:
            with self._session.begin_nested():
                self._session.add(upload)
        except Exception:
            stored.unlink(missing_ok=True)  # never leave a file without its record
            raise
        return UploadResult(upload, message)

    def list(self, student_id: int) -> list[ResumeUpload]:
        return list(self._session.scalars(
            select(ResumeUpload)
            .where(ResumeUpload.student_id == student_id)
            .order_by(ResumeUpload.upload_id.desc())
        ))

    def get(self, student_id: int, upload_id: int) -> ResumeUpload | None:
        return self._session.scalars(select(ResumeUpload).where(
            ResumeUpload.upload_id == upload_id, ResumeUpload.student_id == student_id
        )).first()

    def purge_expired(self) -> int:
        """Delete every file past its retention date (REQ-F1-18). Returns how many."""
        now = self._now()
        due = self._session.scalars(select(ResumeUpload).where(
            ResumeUpload.deleted_at.is_(None), ResumeUpload.delete_after <= now
        )).all()
        for upload in due:
            Path(upload.stored_path).unlink(missing_ok=True)
            upload.deleted_at = now
        self._session.flush()
        return len(due)
