"""UP-6 resume upload against real Postgres (see backend/conftest.py)."""

import asyncio
import threading
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from backend.database import get_session
from backend.main import app
from backend.profile import resume_cleanup
from backend.profile.models import ResumeUpload
from backend.profile.resume import ResumeService
from backend.profile.tests.resume_files import (
    RESUME_LINES,
    docx_with_text,
    docx_zip_bomb,
    encrypted_pdf,
    pdf_with_text,
    pdf_without_text,
)

URL = "/api/v1/profile/resume"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def storage(tmp_path, monkeypatch) -> Path:
    monkeypatch.setenv("RESUME_STORAGE_DIR", str(tmp_path))
    return tmp_path


@pytest.fixture
def client(db_session, storage):
    app.dependency_overrides[get_session] = lambda: db_session
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def student(client, make_user, bearer):
    headers = bearer(make_user())
    assert client.post("/api/v1/profile", json={"name": "Asha"}, headers=headers).status_code == 201
    return headers


def upload(client, headers, name, data, content_type=PDF):
    return client.post(URL, files={"file": (name, data, content_type)}, headers=headers)


def rows(db_session) -> int:
    return db_session.scalar(select(func.count()).select_from(ResumeUpload))


def stored_files(storage: Path) -> list[Path]:
    return [p for p in storage.iterdir() if p.is_file()]


# --- Accepted files (REQ-F1-08, F1-10) -------------------------------------------


@pytest.mark.parametrize(
    ("name", "data", "content_type", "file_type"),
    [
        ("asha-resume.pdf", pdf_with_text(RESUME_LINES), PDF, "pdf"),
        ("asha-resume.docx", docx_with_text(RESUME_LINES), DOCX, "docx"),
    ],
)
def test_readable_resume_is_stored_and_extracted(
    client, storage, student, name, data, content_type, file_type
):
    response = upload(client, student, name, data, content_type)

    assert response.status_code == 201
    body = response.json()
    assert (body["status"], body["file_type"], body["file_name"]) == ("extracted", file_type, name)
    assert body["size_bytes"] == len(data)
    assert body["message"] is None and body["manual_entry"] is None
    assert body["deleted"] is False
    files = stored_files(storage)
    assert len(files) == 1 and files[0].read_bytes() == data
    assert files[0].name != name  # stored under a random name


def test_file_kept_for_30_days_after_processing(client, db_session, student):
    upload(client, student, "cv.pdf", pdf_with_text(RESUME_LINES))
    row = db_session.scalar(select(ResumeUpload))
    assert row.delete_after - row.processed_at == timedelta(days=30)


# --- Rejected files (REQ-F1-08, F1-09; TC-F1-05, TC-F1-06) ----------------------


def test_tc_f1_05_text_file_is_rejected_with_the_exact_message(
    client, db_session, storage, student
):
    text_file = b"Skills: Python, SQL" * 600  # about 12 KB, as in the test plan
    response = upload(client, student, "resume_sample.txt", text_file, "text/plain")
    assert response.status_code == 415
    assert response.json() == {"detail": "Only PDF and DOCX files are accepted."}
    assert rows(db_session) == 0 and stored_files(storage) == []


def test_tc_f1_06_file_over_5_mb_is_rejected_with_the_exact_message(
    client, db_session, storage, student
):
    large = pdf_with_text(RESUME_LINES) + b"%" * int(6.2 * 1024 * 1024)
    response = upload(client, student, "large_resume.pdf", large)
    assert response.status_code == 413
    assert response.json() == {"detail": "File exceeds the 5 MB limit."}
    assert rows(db_session) == 0 and stored_files(storage) == []


@pytest.mark.parametrize(
    ("name", "data", "content_type"),
    [
        ("resume.pdf", b"Skills: Python", PDF),  # PDF name and type, not a PDF
        ("photo.png", b"\x89PNG\r\n\x1a\n" + b"\0" * 64, "image/png"),
        ("archive.docx", b"PK\x03\x04 not really a zip", DOCX),
        ("empty.pdf", b"", PDF),
    ],
)
def test_type_is_decided_by_content_not_name(client, db_session, student, name, data, content_type):
    response = upload(client, student, name, data, content_type)
    assert response.status_code == 415
    assert rows(db_session) == 0


def test_real_docx_with_a_pdf_name_is_accepted_as_docx(client, student):
    response = upload(client, student, "resume.pdf", docx_with_text(RESUME_LINES), PDF)
    assert response.status_code == 201
    assert response.json()["file_type"] == "docx"


# --- Unreadable and failed (REQ-F1-11, F1-38; TC-F1-07) ------------------------


def test_tc_f1_07_scanned_resume_offers_manual_entry_and_is_not_kept(
    client, storage, student
):
    response = upload(client, student, "scanned_only.pdf", pdf_without_text())

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "unreadable"
    assert "add your skills manually" in body["message"]
    assert body["manual_entry"] == "/api/v1/profile/skills"
    assert body["deleted"] is True
    assert stored_files(storage) == []


@pytest.mark.parametrize(
    ("name", "data", "content_type"),
    [
        ("locked.pdf", encrypted_pdf(), PDF),
        ("broken.pdf", b"%PDF-1.7\n" + bytes(range(256)) * 40, PDF),
        ("bomb.docx", docx_zip_bomb(), DOCX),
    ],
)
def test_unprocessable_file_fails_cleanly(client, storage, student, name, data, content_type):
    response = upload(client, student, name, data, content_type)

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "failed"
    assert body["message"] == "We could not process this resume. Your profile has not been changed."
    assert body["deleted"] is True
    assert stored_files(storage) == []


def test_profile_is_unchanged_by_a_failed_upload(client, student):
    before = client.get("/api/v1/profile", headers=student).json()
    upload(client, student, "broken.pdf", b"%PDF-1.7\n" + bytes(range(256)) * 40)
    assert client.get("/api/v1/profile", headers=student).json() == before


# --- Names, listing, ownership --------------------------------------------------------


def test_path_in_the_file_name_is_dropped(client, storage, student):
    response = upload(client, student, "../../etc/passwd.pdf", pdf_with_text(RESUME_LINES))
    assert response.json()["file_name"] == "passwd.pdf"
    assert all(p.parent == storage for p in stored_files(storage))


def test_uploads_are_listed_newest_first_and_readable_by_id(client, student):
    first = upload(client, student, "one.pdf", pdf_with_text(RESUME_LINES)).json()
    second = upload(client, student, "two.docx", docx_with_text(RESUME_LINES), DOCX).json()

    listed = client.get(URL, headers=student).json()
    assert [u["upload_id"] for u in listed] == [second["upload_id"], first["upload_id"]]
    one = client.get(f"{URL}/{first['upload_id']}", headers=student).json()
    assert one["file_name"] == "one.pdf"


def test_another_students_upload_is_not_found(client, make_user, bearer, student):
    mine = upload(client, student, "cv.pdf", pdf_with_text(RESUME_LINES)).json()
    other = bearer(make_user())
    client.post("/api/v1/profile", json={"name": "Ravi"}, headers=other)

    assert client.get(URL, headers=other).json() == []
    response = client.get(f"{URL}/{mine['upload_id']}", headers=other)
    assert response.status_code == 404
    assert response.json() == {"detail": "Resume not found"}


def test_uploading_without_a_profile_returns_404(client, make_user, bearer):
    response = upload(client, bearer(make_user()), "cv.pdf", pdf_with_text(RESUME_LINES))
    assert response.status_code == 404


def test_requires_authentication(client):
    assert upload(client, {}, "cv.pdf", pdf_with_text(RESUME_LINES)).status_code == 401


# --- Retention (REQ-F1-18) ---------------------------------------------------------------


def test_files_are_deleted_only_after_30_days(client, db_session, storage, student):
    upload(client, student, "cv.pdf", pdf_with_text(RESUME_LINES))
    row = db_session.scalar(select(ResumeUpload))

    day_29 = ResumeService(db_session, now=lambda: row.processed_at + timedelta(days=29))
    assert day_29.purge_expired() == 0
    assert len(stored_files(storage)) == 1

    day_30 = ResumeService(db_session, now=lambda: row.processed_at + timedelta(days=30))
    assert day_30.purge_expired() == 1
    assert stored_files(storage) == []
    assert row.deleted_at is not None
    assert day_30.purge_expired() == 0  # already handled


def test_cleanup_command_deletes_expired_files(client, db_session, storage, student, monkeypatch):
    upload(client, student, "cv.pdf", pdf_with_text(RESUME_LINES))
    db_session.execute(
        ResumeUpload.__table__.update().values(delete_after=datetime.now(UTC) - timedelta(days=1))
    )
    monkeypatch.setattr(resume_cleanup, "get_session", lambda: iter([db_session]))
    monkeypatch.setattr(db_session, "close", lambda: None)

    assert resume_cleanup.main() == 1
    assert stored_files(storage) == []


def test_cleanup_runs_periodically_and_survives_a_failed_run(monkeypatch):
    runs, sleeps = [], []

    def purge_once():
        runs.append(1)
        if len(runs) == 1:
            raise RuntimeError("database unavailable")
        return 2

    async def sleep(seconds):
        sleeps.append(seconds)
        if len(sleeps) == 3:
            raise asyncio.CancelledError

    monkeypatch.setattr(resume_cleanup, "purge_once", purge_once)
    monkeypatch.setattr(resume_cleanup.asyncio, "sleep", sleep)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(resume_cleanup.purge_periodically(60))
    assert len(runs) == 3  # the failed first run did not stop the next ones
    assert sleeps == [60, 60, 60]


def test_the_api_purges_at_startup_without_any_cron_job(monkeypatch):
    purged = threading.Event()
    monkeypatch.setattr(resume_cleanup, "purge_once", lambda: purged.set() or 0)
    monkeypatch.setenv("RESUME_CLEANUP_INTERVAL_SECONDS", "3600")
    with TestClient(app):
        assert purged.wait(timeout=5)


def test_cleanup_can_be_turned_off(monkeypatch):
    purged = threading.Event()
    monkeypatch.setattr(resume_cleanup, "purge_once", lambda: purged.set() or 0)
    monkeypatch.setenv("RESUME_CLEANUP_INTERVAL_SECONDS", "0")
    with TestClient(app):
        assert not purged.wait(timeout=0.5)


# --- Extraction details and failure safety ------------------------------------------


def test_text_inside_docx_tables_is_extracted():
    """Many resumes lay skills out in a table."""
    import io

    from docx import Document

    from backend.profile.resume import extract_text

    document = Document()
    document.add_paragraph("Asha Rao")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Kubernetes, TypeScript"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), "docx")
    assert "Kubernetes, TypeScript" in text


def test_no_file_is_left_behind_when_the_record_cannot_be_saved(db_session, storage):
    from sqlalchemy.exc import IntegrityError

    with pytest.raises(IntegrityError):  # student 999999 does not exist
        ResumeService(db_session).upload(999_999, "cv.pdf", pdf_with_text(RESUME_LINES))
    assert stored_files(storage) == []
